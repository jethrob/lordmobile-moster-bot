"""Find new LordsBot xlsx exports, parse them, and store them in the database."""
import asyncio
import datetime as dt
import logging
import re
import time
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .db import GUILD_LIST_COLUMNS, HUNT_COLUMNS, MANUAL_IMPORT, Database

log = logging.getLogger(__name__)

MIN_FILE_AGE_SECONDS = 30  # skip files LordsBot may still be writing
REPORT_MAX_AGE_DAYS = 2  # backfilled old files are imported silently, only fresh ones get a report

# Spreadsheet column index for each field (the blank spacer columns are skipped).
HUNT_INDEXES = [0, 1, 2, 3, 4, 6, 7, 8, 9, 10, 12, 13, 14, 15, 16, 18, 19, 21, 22, 24, 25]
GUILD_LIST_INDEXES = list(range(10))
KINDS = {  # first header cell -> (kind, columns, indexes)
    "User ID": ("hunts", HUNT_COLUMNS, HUNT_INDEXES),
    "IGG ID": ("guild_list", GUILD_LIST_COLUMNS, GUILD_LIST_INDEXES),
}


class UnknownFileError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedFile:
    kind: str
    day: dt.date
    rows: list[dict]


def parse_file(path: Path) -> ParsedFile:
    """Parse a GIFT_STATS (hunts) or GUILD_LIST export. The type comes from the header row, not the filename."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    if not sheet_rows or sheet_rows[0][0] not in KINDS:
        raise UnknownFileError(f"not a GIFT_STATS or GUILD_LIST export: {path.name}")
    kind, columns, indexes = KINDS[sheet_rows[0][0]]
    rows = []
    for raw in sheet_rows[1:]:
        raw = tuple(raw) + (None,) * (max(indexes) + 1 - len(raw))  # some rows are shorter than the header
        # User ID 0 marks the "Total" and "Anonymous" summary rows.
        if not isinstance(raw[0], int) or raw[0] == 0 or raw[1] is None:
            continue
        rows.append({c: _clean(raw[i]) for c, i in zip(columns, indexes)})
    return ParsedFile(kind, file_day(path), rows)


def file_day(path: Path) -> dt.date:
    """Date from names like '2024-08-30 00.00 GIFT_STATS -R-.xlsx' or 'Guild-2024-28-8 00-00 -R-.xlsx'.

    LordsBot writes either Y-M-D or Y-D-M depending on locale; when both readings are valid dates the one
    closest to the file's modified time wins.
    """
    modified = dt.date.fromtimestamp(path.stat().st_mtime)
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", path.name)
    if not match:
        return modified
    year, a, b = map(int, match.groups())
    candidates = [d for d in (_date(year, a, b), _date(year, b, a)) if d]
    return min(candidates, key=lambda d: abs(d - modified)) if candidates else modified


def _date(year, month, day):
    try:
        return dt.date(year, month, day)
    except ValueError:
        return None


def _clean(value):
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time, dt.timedelta)):  # some cells hold only a time of day
        return str(value)
    return None if value == "" else value


def discover_castles(root: Path) -> dict[str, Path]:
    """Castle IGG ID -> export folder.

    LordsBot keeps one folder per castle: C:\\LordsBot\\config\\<IGG>\\stats\\exported. `root` may be the
    config folder itself or the LordsBot folder above it.
    """
    found = {}
    for pattern in ("*/stats/exported", "config/*/stats/exported"):
        for folder in root.glob(pattern):
            igg = folder.parent.parent.name
            if igg.isdigit() and folder.is_dir():
                found.setdefault(igg, folder)
    return dict(sorted(found.items(), key=lambda item: int(item[0])))


def discover_folders(root: Path) -> list[Path]:
    return list(discover_castles(root).values())


def pending_files(db: Database, link: dict, now: float | None = None) -> list[Path]:
    folder = Path(link["folder"])
    if not folder.is_dir():
        return []
    cutoff = (now or time.time()) - MIN_FILE_AGE_SECONDS
    files = [p for p in folder.glob("*.xlsx") if not p.name.startswith("~$") and p.stat().st_mtime < cutoff]
    return sorted((p for p in files if not db.is_imported(p, link["id"])), key=lambda p: p.stat().st_mtime)


def import_file(db: Database, link: dict, path: Path) -> ParsedFile | None:
    """Import one file for one link (or a manual import, see import_folder).

    Failures are recorded (shown on the Activity page) instead of raised.
    """
    link_id, guild = link["id"], link["discord_guild_id"]
    try:
        parsed = parse_file(path)
        db.upsert_rows(parsed.kind, guild, parsed.day, parsed.rows)
        db.record_import(path, link_id, guild, parsed.kind, parsed.day.isoformat(), len(parsed.rows))
        log.info("Imported %s (%s, %s rows) for link %s", path, parsed.kind, len(parsed.rows), link["id"])
        return parsed
    except UnknownFileError as e:
        db.record_import(path, link_id, guild, error=str(e))
        log.info("Skipped %s: %s", path, e)
    except Exception as e:  # locked file, corrupt workbook, ... -> visible on Activity page, re-import from there
        db.record_import(path, link_id, guild, error=f"{type(e).__name__}: {e}")
        log.exception("Import failed: %s", path)
    return None


def import_folder(db: Database, folder: Path, guild_id: int, recursive: bool = False) -> dict:
    """One-off import of every export in a folder (e.g. old exports kept elsewhere). Never posts reports.

    Files are always re-read, so this also repairs files that failed before.
    """
    pattern = "**/*.xlsx" if recursive else "*.xlsx"
    counts = {"imported": 0, "skipped": 0, "failed": 0}
    manual = {"id": MANUAL_IMPORT, "discord_guild_id": guild_id}
    for path in sorted(p for p in folder.glob(pattern) if not p.name.startswith("~$")):
        parsed = import_file(db, manual, path)
        error = (db.imported_file(path, MANUAL_IMPORT) or {}).get("error") or ""
        counts["imported" if parsed else "skipped" if error.startswith("not a GIFT_STATS") else "failed"] += 1
    log.info("Folder import of %s for server %s: %s", folder, guild_id, counts)
    return counts


def should_report(link: dict, parsed: ParsedFile, today: dt.date | None = None) -> bool:
    today = today or dt.date.today()
    return bool(link["post_report"] and link["channel_id"]) and (today - parsed.day).days <= REPORT_MAX_AGE_DAYS


async def run_forever(db: Database, on_imported, interval: float = 60):
    """Poll every link's folder. `on_imported(link, parsed)` is awaited for fresh files that want a report."""
    while True:
        try:
            for link in db.links(enabled_only=True):
                for path in await asyncio.to_thread(pending_files, db, link):
                    parsed = await asyncio.to_thread(import_file, db, link, path)
                    if parsed and should_report(link, parsed):
                        try:
                            await on_imported(link, parsed)
                        except Exception:
                            log.exception("Posting report failed for %s", path)
        except Exception:
            log.exception("Importer loop error")
        await asyncio.sleep(interval)
